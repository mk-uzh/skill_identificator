"""Agreement between the gold standard (your coding in the xlsx, columns 'YOUR ...') and a model run.

Usage: python gld_02_agreement.py --gold coding_proposal_n120_seed42.xlsx --run runs/x.jsonl --split dev

Unit of comparison: main category (first) and its level. Units without an observation count as
category '(none)' / level 'n/o'. Cohen's kappa for category, linearly weighted kappa for level
(n/o is excluded from the level kappa: only units where both sides give a level are compared).
Fix the acceptance threshold BEFORE the test split is run.
"""
import argparse, json
from collections import Counter
import openpyxl

LV = ["l1", "l2", "l3", "l4", "l5"]


def kappa(a, b, labels, weights=None):
    n = len(a)
    if n == 0: return float("nan")
    idx = {l: i for i, l in enumerate(labels)}; k = len(labels)
    O = [[0] * k for _ in range(k)]
    for x, y in zip(a, b): O[idx[x]][idx[y]] += 1
    ra = [sum(r) for r in O]; cb = [sum(O[i][j] for i in range(k)) for j in range(k)]
    w = (lambda i, j: abs(i - j) / (k - 1)) if weights == "linear" else (lambda i, j: 0 if i == j else 1)
    num = sum(w(i, j) * O[i][j] for i in range(k) for j in range(k)) / n
    den = sum(w(i, j) * ra[i] * cb[j] for i in range(k) for j in range(k)) / n ** 2
    return 1 - num / den if den else float("nan")


def load_gold(path, split):
    ws = openpyxl.load_workbook(path, data_only=True)["Units"]
    h = [c.value for c in ws[1]]; ix = {n: i for i, n in enumerate(h)}
    g = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[ix["split"]] != split and split != "all": continue
        cat, lvl = r[ix["YOUR category 1"]], r[ix["YOUR level 1"]]
        if not cat: continue  # not yet judged by the researcher
        g[r[0]] = ("(none)" if cat == "(none)" else cat, lvl or "n/o")
    return g


def load_run(path):
    d = {}
    for line in open(path, encoding="utf-8"):
        o = json.loads(line)
        if "error" in o: continue
        obs = o["observations"]
        d[o["unit"]] = (obs[0]["category"], obs[0]["level"]) if obs else ("(none)", "n/o")
    return d


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--gold", required=True); ap.add_argument("--run", required=True)
    ap.add_argument("--split", default="dev")
    a = ap.parse_args()
    g, m = load_gold(a.gold, a.split), load_run(a.run)
    ids = sorted(set(g) & set(m))
    print(f"gold units judged: {len(g)}, model units: {len(m)}, compared: {len(ids)}")
    if not ids:
        print("Nothing to compare. Fill the 'YOUR category 1' / 'YOUR level 1' columns first."); return
    gc = [g[i][0] for i in ids]; mc = [m[i][0] for i in ids]
    cats = sorted(set(gc) | set(mc))
    print(f"category: raw agreement {sum(x == y for x, y in zip(gc, mc)) / len(ids):.2f}, Cohen's kappa {kappa(gc, mc, cats):.2f}")
    both = [i for i in ids if g[i][1] in LV and m[i][1] in LV]
    if both:
        gl = [g[i][1] for i in both]; ml = [m[i][1] for i in both]
        print(f"level (n={len(both)}): raw agreement {sum(x == y for x, y in zip(gl, ml)) / len(both):.2f}, "
              f"weighted kappa {kappa(gl, ml, LV, 'linear'):.2f}")
    print("most frequent disagreements (gold -> model):")
    for (x, y), c in Counter((g[i][0], m[i][0]) for i in ids if g[i][0] != m[i][0]).most_common(8):
        print(f"  {x} -> {y}: {c}")


if __name__ == "__main__":
    main()
